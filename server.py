# -*- coding: utf-8 -*-
"""高数课件学习助手 server.py（轻量恢复版）

目标：让现有 index.html 可正常访问、上传 PPTX、查看课件结构，并继续使用
m3_generator.py 的“用户主题 + PPT 素材驱动”出题逻辑。

说明：
- 使用标准库 http.server，无 Flask 依赖；
- 上传解析优先用 python-pptx 提取标题/正文/例题；
- MathType/WMF/OCR 精细公式还原不在本轻量恢复版中，已有 slides.json 仍可读取。
"""
import json
import os
import re
import sys
import time
import uuid
import hashlib
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from urllib.parse import urlparse, parse_qs

WORKSPACE = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(WORKSPACE, 'uploads')
sys.path.insert(0, WORKSPACE)

import m3_generator as m3

_SLIDES_CACHE = None
_CURRENT_SOURCE = 'slides.json'


def _slides_path():
    return os.path.join(WORKSPACE, 'slides.json')


def _load_json(path, default):
    try:
        with open(path, encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return default


def _save_json(path, data):
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _slides(force=False):
    global _SLIDES_CACHE
    if force or _SLIDES_CACHE is None:
        path = _slides_path()
        if os.path.exists(path):
            _SLIDES_CACHE = _load_json(path, {'slides': []})
        else:
            _SLIDES_CACHE = {'source': 'none', 'page_count': 0, 'slides': []}
    return _SLIDES_CACHE


def _formula_count(slides):
    return sum(len(p.get('formula_objects') or []) for p in slides)


def _dual_count(slides):
    n = 0
    for p in slides:
        for fo in (p.get('formula_objects') or []):
            if fo.get('text_confidence') == 'dual_channel' or fo.get('ocr_text'):
                n += 1
    return n


def _guess_title(shapes_text):
    for s in shapes_text:
        line = (s or '').strip().split('\n')[0].strip()
        if line:
            return line[:80]
    return ''


def _parse_pptx(path, source_name):
    """轻量 PPTX 解析：提取每页文本/标题/例题。公式对象暂不臆造。"""
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            txt = ''
            try:
                if hasattr(shape, 'text'):
                    txt = shape.text or ''
            except Exception:
                txt = ''
            txt = txt.strip()
            if txt:
                texts.append(txt)
        title = ''
        try:
            if slide.shapes.title and getattr(slide.shapes.title, 'text', None):
                title = slide.shapes.title.text.strip()
        except Exception:
            title = ''
        if not title:
            title = _guess_title(texts)
        body_lines = []
        for txt in texts:
            if txt.strip() and txt.strip() != title:
                body_lines.append(txt.strip())
        body = '\n'.join(body_lines)
        examples = []
        for line in re.split(r'[\r\n]+', body):
            line = line.strip()
            if re.search(r'(例\s*\d*|例题|解[:：]?)', line):
                examples.append(line[:120])
        slides.append({
            'page_no': idx,
            'title': title,
            'body_text': body,
            'formulas': [],
            'formula_objects': [],
            'examples': examples,
            '_ole_count': 0,
            '_wmf_count': 0,
        })
    return {'source': source_name, 'page_count': len(slides), 'slides': slides}


def _extract_multipart_file(handler):
    """从 multipart/form-data 中取第一个文件，返回 (filename, bytes)。"""
    ctype = handler.headers.get('Content-Type', '')
    length = int(handler.headers.get('Content-Length', '0') or '0')
    raw = handler.rfile.read(length) if length else b''
    m = re.search(r'boundary=(?P<b>[^;]+)', ctype)
    if not m:
        return 'upload.pptx', raw
    boundary = m.group('b').strip().strip('"').encode()
    parts = raw.split(b'--' + boundary)
    for part in parts:
        if b'Content-Disposition:' not in part:
            continue
        head, sep, body = part.partition(b'\r\n\r\n')
        if not sep:
            continue
        hm = re.search(br'filename="([^"]*)"', head)
        filename = hm.group(1).decode('utf-8', errors='ignore') if hm else 'upload.pptx'
        body = body.rstrip(b'\r\n')
        if body.endswith(b'--'):
            body = body[:-2]
        return filename or 'upload.pptx', body
    return 'upload.pptx', b''


_TONGJI_CHAPTERS = [
    ('上册', '第一章 函数与极限', ['函数基础','数列极限','函数极限','极限运算法则','夹逼准则','单调有界准则','两个重要极限','无穷小与无穷大','等价无穷小','连续性','间断点分类','闭区间连续函数性质']),
    ('上册', '第二章 导数与微分', ['导数概念','求导法则','复合函数求导','隐函数求导','参数方程求导','高阶导数','微分']),
    ('上册', '第三章 微分中值定理与导数的应用', ['罗尔定理','拉格朗日中值定理','柯西中值定理','洛必达法则','泰勒公式','单调性','凹凸性','拐点','极值与最值','渐近线','曲率']),
    ('上册', '第四章 不定积分', ['原函数','不定积分','基本积分表','第一类换元法','第二类换元法','分部积分法','有理函数积分','三角有理式积分']),
    ('上册', '第五章 定积分', ['定积分概念','定积分性质','牛顿-莱布尼茨公式','积分上限函数','定积分换元法','定积分分部积分法','反常积分']),
    ('上册', '第六章 定积分的应用', ['平面图形面积','旋转体体积','弧长','变力做功','水压力','引力']),
    ('上册', '第七章 微分方程', ['微分方程基本概念','可分离变量方程','齐次方程','一阶线性微分方程','可降阶高阶方程','二阶常系数线性微分方程']),
    ('下册', '第八章 向量代数与空间解析几何', ['向量线性运算','数量积','向量积','混合积','平面方程','空间直线方程','夹角与距离','柱面','旋转曲面','二次曲面','空间曲线']),
    ('下册', '第九章 多元函数微分法及其应用', ['多元函数极限','多元函数连续','偏导数','全微分','多元复合函数求导','隐函数求导','切线与法平面','切平面与法线','方向导数','梯度','多元函数极值','条件极值','拉格朗日乘数法']),
    ('下册', '第十章 重积分', ['二重积分','极坐标二重积分','三重积分','柱面坐标','球面坐标','重积分应用']),
    ('下册', '第十一章 曲线积分与曲面积分', ['第一类曲线积分','第二类曲线积分','格林公式','路径无关','第一类曲面积分','第二类曲面积分','高斯公式','斯托克斯公式']),
    ('下册', '第十二章 无穷级数', ['常数项级数','正项级数审敛法','比较审敛法','比值审敛法','根值审敛法','交错级数','绝对收敛','条件收敛','幂级数','收敛半径','收敛域','泰勒级数','麦克劳林级数','傅里叶级数']),
]

_TOPIC_ALIASES = [
    ('第一类换元法', ['第一类换元','凑微分','2x','x^2']),
    ('第二类换元法', ['第二类换元','三角代换','根式代换']),
    ('不定积分', ['不定积分','原函数']),
    ('定积分换元法', ['定积分','上下限','积分上限']),
    ('等价无穷小', ['等价无穷小','无穷小']),
    ('两个重要极限', ['sin x','重要极限','(1+1/x)^x']),
    ('函数极限', ['极限','lim']),
    ('导数概念', ['导数','可导']),
    ('复合函数求导', ['复合函数','链式法则']),
    ('隐函数求导', ['隐函数']),
    ('单调性', ['单调','递增','递减']),
    ('凹凸性', ['凹凸','凸性','拐点']),
    ('洛必达法则', ['洛必达']),
    ('泰勒公式', ['泰勒']),
]


def _standard_chapters():
    return [
        {'name': ch, 'volume': vol, 'from': None, 'to': None, 'desc': '同济高等数学标准目录', 'kids': topics}
        for vol, ch, topics in _TONGJI_CHAPTERS
    ]


def _all_standard_topics():
    out = []
    for vol, ch, topics in _TONGJI_CHAPTERS:
        for t in topics:
            out.append({'name': t, 'chapter': ch, 'volume': vol})
    return out


def _match_standard_topic(text):
    txt = str(text or '')
    for name, keys in _TOPIC_ALIASES:
        if any(k in txt for k in keys):
            return name
    for item in _all_standard_topics():
        if item['name'] in txt:
            return item['name']
    return None


def _topic_name(title, body):
    text = (title + '\n' + body).strip()
    rules = [
        ('曲线的凹凸性', ('凹凸', '凸性')),
        ('拐点', ('拐点',)),
        ('函数的单调性', ('单调', '增减')),
        ('换元积分', ('换元', '凑微分', '代换')),
        ('积分', ('积分',)),
        ('导数', ('导数', '求导')),
        ('极限', ('极限',)),
    ]
    for name, keys in rules:
        if any(k in text for k in keys):
            return name
    base = title or body[:30] or ''
    base = re.sub(r'例\s*\d*.*$', '例题', base).strip()
    return base[:18] + ('...' if len(base) > 18 else '')


def _knowledge_nodes():
    data = _slides()
    nodes = []
    for p in data.get('slides', []):
        title = (p.get('title') or '').strip()
        body = (p.get('body_text') or '').strip()
        fos = p.get('formula_objects') or []
        examples = p.get('examples') or []
        if not (title or body or fos or examples):
            continue
        name = _topic_name(title, body) or f"P{p.get('page_no')}"
        first_formula = ''
        for fo in fos:
            if (fo.get('text') or '').strip():
                first_formula = fo.get('text').strip()
                break
        nodes.append({
            'page_no': p.get('page_no'),
            'name': name,
            'source': 'title' if title else ('formula' if first_formula else 'body'),
            'formula': first_formula,
            'n_formulas': len(fos),
            'examples': examples[:3],
        })
    return {'nodes': nodes, 'total_pages': len(data.get('slides', [])), 'standard_topics': _all_standard_topics()}


def _syllabus():
    data = _slides()
    slides = data.get('slides', [])
    current_hits = []
    seen = set()
    for p in slides:
        text = ((p.get('title') or '') + ' ' + (p.get('body_text') or ''))
        hit = _match_standard_topic(text)
        if hit and hit not in seen:
            seen.add(hit)
            current_hits.append({'topic': hit, 'page': p.get('page_no')})
    return {
        'chapters': _standard_chapters(),
        'standard_topics': _all_standard_topics(),
        'current_ppt_hits': current_hits,
        'path_hint': '按同济高数标准目录组织知识点；当前 PPT 命中的知识点会用于优先出题。'
    }


def _current():
    data = _slides()
    slides = data.get('slides', [])
    return {
        'loaded': bool(slides),
        'source': data.get('source') or _CURRENT_SOURCE,
        'page_count': len(slides),
        'n_formula_objects': _formula_count(slides),
        'n_dual_channel': _dual_count(slides),
        'cequ': '轻量解析' if _formula_count(slides) == 0 else '已加载',
    }


def api_practice(topic, diff, n):
    try:
        diff = int(diff)
        n = int(n)
    except (ValueError, TypeError):
        diff, n = 2, 5
    slides = _slides().get('slides', [])
    bank = m3.generate_from_slides(slides, topic=topic, difficulty_level=diff, n=n, seed=None)
    not_in_ppt = bool(topic) and not bank
    m3.save_bank(bank, path=os.path.join(WORKSPACE, 'question_bank.json'))
    return {
        'tilted_topic': topic or None,
        'used_constraint': bool(topic),
        'n_generated': len(bank),
        'n_passed': sum(1 for b in bank if b['_validation']['passed']),
        'not_in_ppt': not_in_ppt,
        'difficulty_label': {1: 'L1 基础', 2: 'L2 标准', 3: 'L3 综合'}.get(diff, f'L{diff}'),
        'questions': [
            {'qid': b['qid'], 'difficulty': b['difficulty'], 'level': b['difficulty_level'],
             'prompt': b['prompt'], 'ref': b['ref_answer_str'], 'topic': b.get('topic'),
             'kind': b.get('_kind', 'slide_concept'), 'steps': b.get('steps', []),
             'ppt_pages': b.get('ppt_pages', []),
             'evidence': (b.get('_validation') or {}).get('evidence', ['-'])[0]}
            for b in bank
        ],
    }


def _snippet_for_page(p):
    title = (p.get('title') or '').strip()
    body = (p.get('body_text') or '').strip()
    examples = '；'.join(map(str, (p.get('examples') or [])[:2]))
    formulas = []
    for fo in (p.get('formula_objects') or [])[:6]:
        t = (fo.get('text') or '').strip()
        ot = (fo.get('ocr_text') or '').strip()
        if t:
            formulas.append(t)
        elif ot:
            formulas.append(ot)
    return {
        'page_no': p.get('page_no'),
        'title': title,
        'body': body[:500],
        'examples': examples[:260],
        'formulas': formulas,
        'text': '\n'.join(x for x in [title, body[:500], examples[:260], ' '.join(formulas)] if x),
    }


def _topk_chunks(slides, q, k=5):
    """返回 Top-K 片段。轻量 BM25/词面近似：完整问题命中 + 双字词命中 + 数学关键词命中。"""
    q = (q or '').strip().lower()
    if not q:
        return []
    # 拆 query：中文双字滑窗 + 英文/数字词
    tokens = set(re.findall(r'[A-Za-z0-9]{2,}', q))
    for run in re.findall(r'[\u4e00-\u9fff]+', q):
        if len(run) == 1:
            tokens.add(run)
        else:
            for i in range(len(run)-1):
                tokens.add(run[i:i+2])
    stop = {'什么','怎么','如何','为什么','这个','那个','请问','一下','说明','解释','求解'}
    tokens = {t for t in tokens if t not in stop}
    scored = []
    for p in slides:
        sn = _snippet_for_page(p)
        text = sn['text'].lower()
        if not text:
            continue
        score = 0
        if q and q in text:
            score += 8
        for tk in tokens:
            if tk and tk in text:
                score += 2 if len(tk) >= 2 else 1
        # 标题命中加权
        title = sn['title'].lower()
        for tk in tokens:
            if tk and tk in title:
                score += 2
        if score:
            sn['score'] = score
            scored.append(sn)
    scored.sort(key=lambda x: x.get('score', 0), reverse=True)
    return scored[:k]


def _can_answer_from_chunks(q, chunks):
    if not chunks:
        return False
    # 至少有一个较强命中，或多个弱命中，才认为课件可回答
    best = max(c.get('score', 0) for c in chunks)
    total = sum(c.get('score', 0) for c in chunks)
    return best >= 4 or total >= 8


def _generic_answer(q):
    """无课件直接依据时的通用知识补充。可替换为 Agnes 自身知识回答。"""
    q = q.strip()
    if any(k in q for k in ('凹凸', '凹凸性', '凸性', '拐点')):
        return ('函数的凹凸性描述曲线弯曲方向，通常用二阶导数判断。\n'
                '1. 若在某区间内 f′′(x)>0，曲线通常是凹向上的，也常称为凸。\n'
                '2. 若在某区间内 f′′(x)<0，曲线通常是凹向下的，也常称为凹。\n'
                '3. 凹凸性发生改变的点可能是拐点，但还需要函数在该点附近有定义并满足相应条件。')
    if any(k in q for k in ('换元', '凑微分', '积分')):
        return ('换元积分的核心思想是把被积函数中的复合结构看成新的变量。\n'
                '1. 先观察是否存在内函数 g(x) 及其微分因子 g′(x)dx。\n'
                '2. 令 u=g(x)，把原积分转化为关于 u 的较简单积分。\n'
                '3. 积分完成后要把 u 回代为原来的 x 表达式。')
    if any(k in q for k in ('导数', '求导')):
        return ('导数描述函数的瞬时变化率。\n'
                '1. 它可以理解为函数图像在某点切线的斜率。\n'
                '2. 常用规则包括和差、乘积、商和链式法则。\n'
                '3. 处理复合函数时，关键是外层求导后再乘以内层导数。')
    if any(k in q for k in ('极限',)):
        return ('极限描述变量趋近某个值时函数值的趋势。\n'
                '1. 简单情形可直接代入。\n'
                '2. 未定式常用因式分解、通分、有理化或等价无穷小处理。\n'
                '3. 满足条件时也可使用洛必达法则。')
    return ('可以先明确问题涉及的对象、条件和结论，再选择对应定义或定理。\n'
            '1. 概念题重点说明定义、条件和几何意义。\n'
            '2. 计算题重点写出关键公式和变形步骤。\n'
            '3. 证明题重点说明依据和逻辑链条。')


def _clean_context_text(text):
    """清理课件片段：去掉孤立问号、空白、公式/图片占位导致的断句。"""
    text = (text or '').replace('\r', '\n')
    raw_lines = re.split(r'[\n。；;]+', text)
    lines = []
    for line in raw_lines:
        line = re.sub(r'\s+', ' ', line).strip(' ，,、:：')
        if not line:
            continue
        if line in {'?', '？', '.', '。', '-', '—'}:
            continue
        # 过滤明显的孤立公式/占位符：过短且无中文说明
        if len(line) <= 2 and not re.search(r'[\u4e00-\u9fff]{2,}', line):
            continue
        # 过滤只含符号/单字母的片段
        if re.fullmatch(r'[A-Za-z0-9=+\-*/^()\[\]{}., ]{1,8}', line):
            continue
        lines.append(line)
    # 去重保序
    out = []
    seen = set()
    for line in lines:
        if line not in seen:
            seen.add(line)
            out.append(line)
    return out


def _answer_from_courseware(q, chunks):
    """按指定格式生成课件内答案：一句话概括 + 分点 + 来源页码。"""
    pages = []
    cleaned = []
    titles = []
    for c in chunks[:5]:
        pn = c.get('page_no')
        if pn not in pages:
            pages.append(pn)
        if c.get('title'):
            titles.append(c['title'].strip())
        cleaned.extend(_clean_context_text(c.get('body') or ''))
        cleaned.extend(_clean_context_text(c.get('examples') or ''))

    qtext = q or ''
    # 常见高数专题做结构化概括，避免直接拼原文。
    if any(k in qtext for k in ('凹凸', '凹凸性', '凸性')):
        summary = '凹凸性描述函数图像在区间上的弯曲方向。'
        points = [
            '通常用二阶导数判断：二阶导数为正时图像偏向上弯，为负时偏向下弯。',
            '研究凹凸性时，一般先求二阶导数，再找二阶导数为零或不存在的可疑分界点。',
            '若二阶导数在某点两侧符号发生改变，该点可能对应函数图像的拐点。',
        ]
    elif any(k in qtext for k in ('单调', '增减')):
        summary = '单调性描述函数在区间上随自变量变化的增减趋势。'
        points = [
            '通常用一阶导数判断：导数为正时函数递增，导数为负时函数递减。',
            '求单调区间时，一般先求导，再找导数为零或不存在的点划分区间。',
            '最后在各区间判断导数符号，并据此写出递增区间和递减区间。',
        ]
    elif any(k in qtext for k in ('换元', '凑微分', '积分')):
        summary = '换元积分是把复杂积分转化为简单积分的方法。'
        points = [
            '第一类换元法的关键是识别复合函数结构，并凑出内函数的微分。',
            '设 u=g(x) 后，要把 g′(x)dx 一起替换成 du，使积分变成关于 u 的形式。',
            '积分完成后要回代原变量；定积分换元时还要同步改上下限。',
        ]
    else:
        # 未覆盖专题：只概括能确定的内容，不硬凑。
        usable = cleaned[:3]
        if not usable:
            return None, pages
        summary = '课件相关片段可提供部分线索。'
        points = []
        for item in usable:
            # 改写为“课件指出/涉及”，避免原文拼接
            points.append('课件涉及“' + item[:70] + '”这一点，可作为回答该问题的依据。')

    # 如果检索片段过弱，只保留可确定部分。
    if not points:
        return None, pages
    answer = summary[:50]
    answer += '\n'
    for i, pt in enumerate(points[:3], start=1):
        answer += f'{i}. {pt}\n'
    answer += '（来源：' + '、'.join(f'P{p}' for p in pages if p is not None) + '）'
    return answer[:10000], pages


def _agnes_answer_from_context(q, chunks):
    """将“用户问题 + Top-K 片段”交给模型生成答案的封装。

    目标 prompt 规则：
    - 不直接复述片段原文，要用自己的话概括；
    - 忽略空白、孤立问号、未完成句子、公式/图片占位；
    - 格式：第一行一句话概括（≤50字），然后 1. 2. 3. 分点展开；
    - 每点不超过两行，末尾标注来源页码，总字数 ≤10000；
    - 信息不足时只概括能确定的部分，不硬凑。
    """
    if _can_answer_from_chunks(q, chunks):
        ans, pages = _answer_from_courseware(q, chunks)
        if ans:
            return ans, pages, True
    ans = '课件中未直接涉及该内容，以下基于通用知识补充：' + _generic_answer(q) + '\n（此内容非来自上传课件）'
    return ans[:10000], [], False


def _is_courseware_overview_question(q):
    text = str(q or '')
    return any(k in text for k in ('这个课件讲了什么', '课件讲了什么', '主要讲什么', '主要内容', '概括', '总结课件', '课件内容', '这节课讲了什么', '本课讲了什么'))


def _courseware_overview_answer(slides):
    """回答“课件讲了什么”这类总览问题。"""
    if not slides:
        return '当前还没有可概括的课件内容。', []
    pages = []
    hits = []
    for p in slides:
        title = (p.get('title') or '').strip()
        body = (p.get('body_text') or '').strip()
        text = (title + ' ' + body).strip()
        if not text:
            continue
        topic = _match_standard_topic(text)
        if topic and topic not in hits:
            hits.append(topic)
        if p.get('page_no') is not None:
            pages.append(p.get('page_no'))
    # 若标准匹配较少，从章节标题补充
    titles = []
    for p in slides:
        title = (p.get('title') or '').strip()
        if title and not re.fullmatch(r'\d+', title) and title not in titles:
            titles.append(title)
    topics = hits[:6] or titles[:6]
    source_pages = []
    if pages:
        source_pages = sorted(set([pages[0], pages[-1]] + pages[:3]))[:5]
    if topics:
        first = '这份课件主要围绕' + '、'.join(topics[:3]) + '展开。'
    else:
        first = '这份课件主要讲解当前章节的核心概念、公式和例题。'
    lines = [first[:50]]
    if topics:
        lines.append('1. 主要知识点：' + '、'.join(topics[:6]) + '。')
    else:
        lines.append('1. 课件包含若干页正文、公式和例题内容。')
    # 依据当前 PPT 常见结构概括，不编造额外章节
    n_formula = sum(len(p.get('formula_objects') or []) for p in slides)
    n_examples = sum(len(p.get('examples') or []) for p in slides)
    lines.append(f'2. 课件共 {len(slides)} 页，包含概念说明、公式推导或计算示例。')
    if n_examples:
        lines.append(f'3. 其中包含 {n_examples} 个例题/示例，可用于后续练习与批改。')
    elif n_formula:
        lines.append(f'3. 其中识别到 {n_formula} 条公式内容，可用于检索和生成练习。')
    else:
        lines.append('3. 可以根据已解析文本继续做检索问答和针对性练习。')
    if source_pages:
        lines.append('（来源：' + '、'.join(f'P{x}' for x in source_pages) + '）')
    else:
        lines.append('（来源：当前上传课件）')
    return '\n'.join(lines), source_pages


def api_qa(q):
    q = (q or '').strip()
    if not q:
        return {'answer': '请输入问题。', 'source_pages': [], 'used_courseware': False, 'chunks': []}
    slides = _slides().get('slides', [])
    if _is_courseware_overview_question(q):
        answer, pages = _courseware_overview_answer(slides)
        return {'answer': answer, 'source_pages': pages, 'used_courseware': True, 'chunks': []}
    chunks = _topk_chunks(slides, q, k=5)
    answer, pages, used = _agnes_answer_from_context(q, chunks)
    # 前端只显示 answer；source_pages 保留给调试/未来扩展，不单独渲染
    return {
        'answer': answer,
        'source_pages': pages,
        'used_courseware': used,
        'chunks': [{'page_no': c.get('page_no'), 'score': c.get('score')} for c in chunks],
    }


def _norm_answer(x):
    x = str(x or '').strip().lower()
    x = x.replace('＋', '+').replace('－', '-').replace('×', '*').replace('÷', '/')
    x = re.sub(r'\s+', '', x)
    x = x.replace('$', '').replace('，', ',').replace('。', '.')
    return x


def _judge_text_answer(stu, ref):
    """文本题兜底判断。判断题要先精确识别，避免“对”匹配到“不对”。"""
    ns, nr = _norm_answer(stu), _norm_answer(ref)
    def polar(v):
        vv = v.replace('。','').replace('.','')
        neg = ['不对','错误','错','否','不正确','false','×','不是']
        pos = ['对','正确','是','true','√']
        for k in neg:
            if vv == k or vv.startswith(k):
                return False
        for k in pos:
            if vv == k or vv.startswith(k):
                return True
        return None
    ps, pr = polar(ns), polar(nr)
    if ps is not None or pr is not None:
        if ps is None or pr is None:
            return False, '判断题需明确回答“对/不对”或“正确/错误”。'
        return ps == pr, '判断题按正误语义精确匹配。'
    if ns == nr:
        return True, '文本完全一致。'
    # 只允许较长文本的包含匹配，避免“对”命中“不对”等短词误判
    if len(ns) >= 4 and (ns in nr or nr in ns):
        return True, '文本近似匹配。'
    return False, '文本答案与参考答案不一致。'


def _try_math_equal(a, b, allow_constant=False):
    """轻量数学等价判断；SymPy 不可用时返回 None。"""
    try:
        if re.search(r'[\u4e00-\u9fff]', str(a or '') + str(b or '')):
            return None, '文本答案不进行符号等价判断。'
        import sympy as sp
        x = sp.symbols('x')
        loc = {'x': x, 'e': sp.E, 'pi': sp.pi, 'ln': sp.log, 'log': sp.log, 'sin': sp.sin, 'cos': sp.cos, 'tan': sp.tan, 'sqrt': sp.sqrt, 'exp': sp.exp}
        aa = str(a or '').replace('^', '**').replace('C', '').replace('+c', '').replace('+C', '')
        bb = str(b or '').replace('^', '**').replace('C', '').replace('+c', '').replace('+C', '')
        ea = sp.sympify(aa, locals=loc)
        eb = sp.sympify(bb, locals=loc)
        diff = sp.simplify(ea - eb)
        if diff == 0:
            return True, 'SymPy 化简后两式差为 0'
        # 不定积分答案允许差常数；其他题型不能把 0 和 1 这种常数差误判为等价。
        if allow_constant and sp.simplify(sp.diff(diff, x)) == 0:
            return True, '两式相差常数，按不定积分等价处理'
        return False, f'SymPy 化简差值为 {diff}'
    except Exception:
        return None, '未能进行符号等价判断，已使用文本近似批改'


def _classify_problem(prompt, topic):
    text = (str(prompt or '') + ' ' + str(topic or '')).lower()
    if any(k in text for k in ['极限', 'lim', 'limit']):
        return 'limit'
    if any(k in text for k in ['导数', '求导', "f'", 'derivative']):
        return 'derivative'
    if any(k in text for k in ['积分', '∫', 'integral']):
        return 'integral'
    if any(k in text for k in ['单调', '递增', '递减']):
        return 'monotonicity'
    if any(k in text for k in ['凹凸', '凸性', '拐点']):
        return 'concavity'
    return 'unknown'


def _sympify_expr(expr):
    import sympy as sp
    x = sp.symbols('x')
    loc = {'x': x, 'e': sp.E, 'E': sp.E, 'pi': sp.pi, 'ln': sp.log, 'log': sp.log,
           'sin': sp.sin, 'cos': sp.cos, 'tan': sp.tan, 'sqrt': sp.sqrt, 'exp': sp.exp}
    s = str(expr or '').strip()
    s = s.replace('^', '**').replace('·', '*').replace('×', '*').replace('÷', '/')
    # 去掉常见 LaTeX 包装
    s = s.replace('$', '').replace('\\', '')
    return sp.sympify(s, locals=loc), x


def _extract_limit_point(prompt):
    text = str(prompt or '')
    # lim(x→1), x->1, x→0, limit(...,x,1)
    m = re.search(r'x\s*[→\-]>?\s*([\-\d\.]+)', text)
    if m:
        try: return float(m.group(1)) if '.' in m.group(1) else int(m.group(1))
        except Exception: pass
    m = re.search(r'limit\([^,]+,\s*x\s*,\s*([\-\d\.]+)\s*\)', text, re.I)
    if m:
        try: return float(m.group(1)) if '.' in m.group(1) else int(m.group(1))
        except Exception: pass
    return None


def _grade_limit_process(prompt, ref, steps_in):
    import sympy as sp
    x = sp.symbols('x')
    point = _extract_limit_point(prompt)
    if point is None:
        point = 0 if 'x→0' in str(prompt) or 'x->0' in str(prompt) else None
    steps = []
    first_wrong = None
    error_type = None
    evidence = ''

    # 最终答案判断
    final = steps_in[-1] if steps_in else ''
    final_eq, final_ev = _try_math_equal(final, ref, allow_constant=('积分' in str(prompt) or 'integral' in str(prompt).lower() or 'integrate' in str(prompt).lower()))
    final_ok = bool(final_eq) if final_eq is not None else (_norm_answer(final) == _norm_answer(ref))

    prev_expr = None
    prev_limit = None
    ref_value = None
    try:
        ref_value, _ = _sympify_expr(ref)
    except Exception:
        ref_value = None

    for i, raw in enumerate(steps_in, start=1):
        expr = str(raw).strip()
        ok = True
        ev = ''
        # 0/0、∞/∞ 这种不能作为等价变形结果
        compact = _norm_answer(expr)
        if compact in {'0/0', '∞/∞', 'inf/inf', 'oo/oo'}:
            ok = False
            ev = '0/0 是未定式，不能作为合法等价步骤或极限值。'
            if first_wrong is None:
                first_wrong, error_type, evidence = f'第{i}步', '未定式误判', ev
        else:
            try:
                e, sx = _sympify_expr(expr)
                if point is not None:
                    lim_val = sp.limit(e, sx, point)
                    # 若上一可解析表达式已有极限值，本步必须保持同一极限
                    if prev_limit is not None and sp.simplify(lim_val - prev_limit) != 0:
                        ok = False
                        ev = f'该步表达式在 x→{point} 时极限为 {lim_val}，与上一步极限 {prev_limit} 不一致。'
                        if first_wrong is None:
                            first_wrong, error_type, evidence = f'第{i}步', '等价变形错误', ev
                    else:
                        ev = f'该步在 x→{point} 时极限为 {lim_val}'
                    prev_limit = lim_val if ok else prev_limit
                else:
                    ev = '未识别极限趋近点，仅记录该步。'
                prev_expr = e
            except Exception as ex:
                # 最后答案可以是数字/文本；中间不可解析则不判正确
                if i < len(steps_in):
                    ok = False
                    ev = f'该中间步骤无法解析为可验证表达式：{ex}'
                    if first_wrong is None:
                        first_wrong, error_type, evidence = f'第{i}步', '步骤无法验证', ev
                else:
                    ev = final_ev
        steps.append({'label': f'第{i}步', 'expr': expr, 'ok': ok, 'evidence': ev})

    process_ok = first_wrong is None
    overall_ok = final_ok and process_ok
    if overall_ok:
        verdict = '正确'
        evidence = final_ev or '最终答案与参考答案一致，过程未发现非法步骤。'
    elif final_ok and not process_ok:
        verdict = '最终答案正确，但过程有错误'
        if not evidence:
            evidence = '最终答案正确，但过程存在未验证或不合法步骤。'
    else:
        verdict = '需订正'
        if first_wrong is None:
            first_wrong = '最终答案'
            error_type = '最终答案错误'
            evidence = final_ev or '最终答案与参考答案不一致。'
    return verdict, overall_ok, steps, first_wrong, error_type, evidence




def _pre_math_text(text):
    s = str(text or '')
    trans = str.maketrans({'²':'**2','³':'**3','⁴':'**4','⁵':'**5','π':'pi','∞':'oo','－':'-','＋':'+','×':'*','÷':'/','·':'*','^':'**'})
    return s.translate(trans)


def _auto_solve_problem(prompt, problem_type):
    """从题目自动求参考答案。当前支持标准/半标准极限、导数、积分题。"""
    try:
        import sympy as sp
        x = sp.symbols('x')
        loc = {'x':x,'e':sp.E,'E':sp.E,'pi':sp.pi,'oo':sp.oo,'ln':sp.log,'log':sp.log,'sin':sp.sin,'cos':sp.cos,'tan':sp.tan,'sqrt':sp.sqrt,'exp':sp.exp}
        raw = _pre_math_text(prompt)
        low = raw.lower()
        # 1) limit(expr, x, a)
        m = re.search(r'limit\((.+?),\s*x\s*,\s*([^\)]+)\)', raw, re.I)
        if m:
            expr = sp.sympify(m.group(1), locals=loc)
            point = sp.sympify(m.group(2), locals=loc)
            val = sp.limit(expr, x, point)
            return str(val), f'自动识别为极限题：limit({sp.sstr(expr)}, x, {point}) = {val}'
        # 2) lim(x→a) expr 或 lim(x->a) expr
        m = re.search(r'lim\s*\(?\s*x\s*(?:→|->|=>|to)\s*([^\)\s，,]+)\s*\)?\s*[:：]?\s*(.+)$', raw, re.I)
        if m:
            point = sp.sympify(m.group(1), locals=loc)
            expr_txt = m.group(2).strip()
            expr_txt = re.sub(r'^[：:，,\s]+', '', expr_txt)
            expr = sp.sympify(expr_txt, locals=loc)
            val = sp.limit(expr, x, point)
            return str(val), f'自动识别为极限题：x→{point} 时 {sp.sstr(expr)} 的极限为 {val}'
        # 3) diff(expr, x) 或 求导
        m = re.search(r'diff\((.+?)(?:,\s*x)?\)', raw, re.I)
        if m:
            expr = sp.sympify(m.group(1), locals=loc)
            val = sp.diff(expr, x)
            return str(val), f'自动识别为导数题：d/dx({sp.sstr(expr)}) = {val}'
        if problem_type == 'derivative' or '求导' in raw or '导数' in raw:
            m = re.search(r'(?:y\s*=|f\(x\)\s*=)?\s*([A-Za-z0-9_+\-*/().\s]+)$', raw)
            if m:
                expr = sp.sympify(m.group(1), locals=loc)
                val = sp.diff(expr, x)
                return str(val), f'自动识别为导数题：d/dx({sp.sstr(expr)}) = {val}'
        # 4) integrate(expr, x) 或 ∫expr dx
        m = re.search(r'integrate\((.+?)(?:,\s*x)?\)', raw, re.I)
        if m:
            expr = sp.sympify(m.group(1), locals=loc)
            val = sp.integrate(expr, x)
            return str(val), f'自动识别为积分题：∫ {sp.sstr(expr)} dx = {val} + C'
        m = re.search(r'∫\s*(.+?)\s*d\s*x', raw)
        if m:
            expr = sp.sympify(m.group(1), locals=loc)
            val = sp.integrate(expr, x)
            return str(val), f'自动识别为积分题：∫ {sp.sstr(expr)} dx = {val} + C'
        return None, '未能自动识别题目，请改写为标准形式，如 limit(sin(x)/x, x, 0)、diff(x**3, x)、integrate(2*x*cos(x**2), x)。'
    except Exception as e:
        return None, f'自动求解失败：{e}。请改写为标准形式，如 limit(sin(x)/x, x, 0)。'


def _record_grade_history(result, prompt=''):
    """批改后回写薄弱档案。同一道题重复提交时更新已有记录，不重复追加。"""
    try:
        topic = result.get('topic') or result.get('problem_type') or '未分类'
        err = result.get('error_type') or ('正确' if result.get('is_correct') else '未分类错误')
        is_correct = bool(result.get('is_correct')) and result.get('verdict') == '正确'
        prompt_norm = re.sub(r'\s+', '', str(prompt or ''))
        # 以题干 + 学生答案/错误表达式作为一次提交的稳定键；同题同答重复提交只更新，不累计次数。
        ans_norm = re.sub(r'\s+', '', str(result.get('student_answer') or result.get('wrong_expr') or ''))
        submit_key = hashlib.sha1((topic + '|' + prompt_norm + '|' + ans_norm).encode('utf-8')).hexdigest()[:16]
        prof_path = os.path.join(WORKSPACE, 'weakness_profile.json')
        hist_path = os.path.join(WORKSPACE, 'weakness_history.json')
        prof = _load_json(prof_path, {})
        if not isinstance(prof, dict):
            prof = {}
        topics = prof.setdefault('topics', {})
        td = topics.setdefault(topic, {'correct': 0, 'total_errors': 0, 'errors': {}, 'recent_rates': [], 'related': [], 'status': 'untouched'})

        related = td.setdefault('related', [])
        old_rec = next((r for r in related if r.get('submit_key') == submit_key), None)
        old_hist = None
        hist = _load_json(hist_path, [])
        if not isinstance(hist, list):
            hist = []
        for h in hist:
            if h.get('submit_key') == submit_key:
                old_hist = h
                break
        old_correct = old_rec.get('correct') if old_rec is not None else (old_hist.get('correct') if old_hist else None)
        old_err = old_rec.get('error_type') if old_rec is not None else (old_hist.get('error_type') if old_hist else None)

        def inc_counter(was_correct, was_err, delta):
            if was_correct is True:
                td['correct'] = max(0, int(td.get('correct', 0)) + delta)
            elif was_correct is False:
                td['total_errors'] = max(0, int(td.get('total_errors', 0)) + delta)
                if was_err:
                    errs = td.setdefault('errors', {})
                    errs[was_err] = max(0, int(errs.get(was_err, 0)) + delta)
                    if errs[was_err] == 0:
                        errs.pop(was_err, None)

        if old_correct is not None:
            inc_counter(old_correct, old_err, -1)
        inc_counter(is_correct, None if is_correct else err, +1)

        new_rate = 100 if is_correct else 0
        rates = td.setdefault('recent_rates', [])
        if old_hist is None and old_rec is None:
            rates.append(new_rate)
        else:
            # 重复提交同一道题：替换最近一次同题记录，不增加趋势长度。
            # 若无法定位具体位置，替换最后一个值。
            if rates:
                rates[-1] = new_rate
            else:
                rates.append(new_rate)
        td['recent_rates'] = rates[-8:]

        now = time.strftime('%Y-%m-%d %H:%M:%S')
        if not is_correct:
            rec = {
                'id': old_rec.get('id') if old_rec else 'G-' + uuid.uuid4().hex[:8],
                'submit_key': submit_key,
                'correct': False,
                'time': now,
                'verdict': result.get('verdict'),
                'error_type': err,
                'prompt': prompt,
                'wrong_expr': result.get('wrong_expr') or result.get('student_answer'),
                'correct_ref': result.get('correct_ref'),
                'diagnosis': result.get('symbol_evidence'),
                'practice': (result.get('correction') or {}).get('practice') if result.get('correction') else '',
            }
            related = [r for r in related if r.get('submit_key') != submit_key]
            related.insert(0, rec)
            td['related'] = related[:20]
        else:
            td['related'] = [r for r in related if r.get('submit_key') != submit_key][:20]

        td['status'] = 'weak' if int(td.get('total_errors', 0)) > int(td.get('correct', 0)) else 'mastered'
        prof['updated_at'] = now
        _save_json(prof_path, prof)

        hist = [h for h in hist if h.get('submit_key') != submit_key]
        hist.insert(0, {'submit_key': submit_key, 'time': now, 'topic': topic, 'correct': is_correct, 'verdict': result.get('verdict'), 'error_type': None if is_correct else err, 'prompt': prompt})
        _save_json(hist_path, hist[:200])
    except Exception:
        pass


def _profile_board():
    prof = _load_json(os.path.join(WORKSPACE, 'weakness_profile.json'), {})
    topics = prof.get('topics', {}) if isinstance(prof, dict) else {}
    diagnosed = len([t for t, d in topics.items() if (d.get('correct',0) or d.get('total_errors',0))])
    weakest_topic, weakest_errors = None, 0
    total_correct = total_errors = 0
    recent = []
    for t, d in topics.items():
        total_correct += int(d.get('correct', 0))
        total_errors += int(d.get('total_errors', 0))
        if int(d.get('total_errors', 0)) > weakest_errors:
            weakest_topic, weakest_errors = t, int(d.get('total_errors', 0))
        recent.extend(d.get('recent_rates', [])[-3:])
    total = total_correct + total_errors
    return {
        'topics': topics,
        'diagnosed_topics': diagnosed,
        'weakest_topic': weakest_topic,
        'weakest_errors': weakest_errors,
        'overall_pct': round(total_correct * 100 / total) if total else None,
        'overall_recent_pct': recent[-8:],
        'overall_recent_avg': round(sum(recent[-8:]) / len(recent[-8:])) if recent[-8:] else None,
    }
def api_grade(payload):
    mode = payload.get('mode', 'A')
    topic = payload.get('topic', '')
    prompt = payload.get('prompt') or payload.get('question') or payload.get('integrand') or ''
    ref = payload.get('ref_answer') or payload.get('ref') or ''
    auto_ref_evidence = ''
    if not ref:
        requested_type_tmp = payload.get('problem_type') or 'auto'
        guessed_type_tmp = requested_type_tmp if requested_type_tmp != 'auto' else _classify_problem(prompt, topic)
        ref, auto_ref_evidence = _auto_solve_problem(prompt, guessed_type_tmp)
        if not ref:
            return {
                'ok': False,
                'mode': mode,
                'problem_type': guessed_type_tmp,
                'verdict': '无法自动批改',
                'is_correct': False,
                'topic': topic,
                'student_answer': payload.get('student_answer') or payload.get('answer') or '',
                'correct_ref': None,
                'first_wrong_step': None,
                'error_type': '题目无法自动识别',
                'symbol_evidence': auto_ref_evidence,
                'correction': {'correction': auto_ref_evidence, 'practice': '请把题目改写为更标准的数学表达式后再提交。'},
                'steps': [],
                'evidence': {'detail': auto_ref_evidence},
            }
    stu = payload.get('student_answer') or payload.get('answer') or ''
    steps_in = payload.get('student_steps') or []
    if isinstance(steps_in, str):
        steps_in = [s.strip() for s in steps_in.split('\n') if s.strip()]
    requested_type = payload.get('problem_type') or 'auto'
    problem_type = requested_type if requested_type != 'auto' else _classify_problem(prompt, topic)

    reliability = '高'
    grade_status = 'correct'
    if mode == 'B' and steps_in and problem_type == 'limit':
        verdict, ok, steps, first_wrong, error_type, evidence = _grade_limit_process(prompt, ref, steps_in)
        stu_final = steps_in[-1]
        if verdict == '最终答案正确，但过程有错误':
            grade_status, reliability = 'partial', '中'
        elif not ok:
            grade_status = 'incorrect'
    else:
        stu_final = steps_in[-1] if (mode == 'B' and steps_in) else stu
        math_eq, evidence = _try_math_equal(stu_final, ref, allow_constant=(problem_type == 'integral'))
        uncertain = False
        if math_eq is None:
            ok, evidence = _judge_text_answer(stu_final, ref)
            # 文字/概念类只在明确正误判断匹配时给高可靠；其他文本不强行判错。
            if '判断题按正误语义精确匹配' not in evidence and problem_type in ('unknown', 'other'):
                uncertain = True
        else:
            ok = bool(math_eq)
        steps = []
        if mode == 'B':
            for i, expr in enumerate(steps_in, start=1):
                is_last = (i == len(steps_in))
                steps.append({
                    'label': f'第{i}步',
                    'expr': expr,
                    'ok': (ok if is_last else None),
                    'evidence': evidence if is_last else '当前题型暂未接入专用过程批改器：该中间步骤未判定对错。'
                })
        first_wrong = None if (ok or uncertain) else ('最终答案' if mode == 'A' else f'第{len(steps_in) or 1}步')
        error_type = None if (ok or uncertain) else '答案不等价或表述不完整'
        if uncertain:
            verdict = '暂不能可靠判定'
            grade_status = 'unreliable'
            reliability = '低'
            ok = False
            evidence = evidence + '；该题属于文本/开放表达，当前不能保证自动判分。'
        elif ok and mode == 'B' and problem_type != 'limit':
            verdict = '部分正确 / 过程未完全验证'
            grade_status = 'partial'
            reliability = '中'
            evidence = evidence + '；最终答案等价，但当前题型暂未进行逐步过程判定。'
        else:
            verdict = '正确' if ok else '需订正'
            grade_status = 'correct' if ok else 'incorrect'

    correction = None if (ok or grade_status == 'unreliable') else {
        'correction': '请对照参考答案检查关键公式、符号和化简步骤。',
        'practice': f'再做一道与“{topic or "当前知识点"}”相关的同类题，并写出完整步骤。'
    }
    result = {
        'ok': True,
        'mode': mode,
        'problem_type': problem_type,
        'verdict': verdict,
        'grade_status': grade_status,
        'reliability': reliability,
        'is_correct': ok and grade_status == 'correct',
        'topic': topic or problem_type,
        'student_answer': stu_final,
        'correct_ref': ref,
        'wrong_expr': None if ok else stu_final,
        'first_wrong_step': first_wrong,
        'error_type': error_type,
        'symbol_evidence': ((auto_ref_evidence + '；') if auto_ref_evidence else '') + evidence,
        'correction': correction,
        'steps': steps,
        'evidence': {'detail': ((auto_ref_evidence + '；') if auto_ref_evidence else '') + evidence},
    }
    _record_grade_history(result, prompt)
    return result



class Handler(SimpleHTTPRequestHandler):
    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path
        qs = parse_qs(parsed.query)
        if path == '/api/practice':
            self._json(api_practice(qs.get('topic', [''])[0].strip(), qs.get('diff', ['2'])[0], qs.get('n', ['5'])[0]))
        elif path == '/api/slides':
            self._json(_slides())
        elif path.startswith('/api/pages/'):
            try:
                no = int(path.split('/')[-1])
            except ValueError:
                self._json({'error': 'bad page_no'}, 400); return
            for p in _slides().get('slides', []):
                if p.get('page_no') == no:
                    self._json(p); return
            self._json({'error': 'page not found'}, 404)
        elif path == '/api/current':
            self._json(_current())
        elif path == '/api/ocr-status':
            self._json({'running': False, 'ready': True, 'result': {'ocr_hit': _dual_count(_slides().get('slides', [])), 'ocr_total': _formula_count(_slides().get('slides', [])), 'secs': 0}})
        elif path == '/api/knowledge':
            self._json(_knowledge_nodes())
        elif path == '/api/kgraph':
            nodes = _knowledge_nodes()['nodes']
            self._json({'hint': '轻量知识图谱已加载。', 'node_status': {str(n['page_no']): {'status': 'untouched', 'matched_topic': None, 'err_stats': {}} for n in nodes}})
        elif path == '/api/syllabus':
            self._json(_syllabus())
        elif path == '/api/search':
            self._json({'ok': True})
        elif path == '/api/qa':
            self._json(api_qa(qs.get('q', [''])[0]))
        elif path == '/api/profile-board':
            self._json(_profile_board())
        elif path.startswith('/api/profile-topic-practice'):
            topic = qs.get('topic', [''])[0]
            r = api_practice(topic, 2, 3)
            self._json(r)
        else:
            super().do_GET()

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path == '/api/upload':
            self._handle_upload()
        elif parsed.path == '/api/grade':
            try:
                length = int(self.headers.get('Content-Length', 0) or 0)
                raw = self.rfile.read(length).decode('utf-8') if length else '{}'
                payload = json.loads(raw or '{}')
            except Exception:
                payload = {}
            self._json(api_grade(payload))
        else:
            self._json({'error': 'unknown POST path', 'path': parsed.path}, 404)

    def _handle_upload(self):
        global _SLIDES_CACHE, _CURRENT_SOURCE
        try:
            os.makedirs(UPLOAD_DIR, exist_ok=True)
            filename, content = _extract_multipart_file(self)
            safe = re.sub(r'[^0-9A-Za-z_.\-\u4e00-\u9fff]+', '_', filename or 'upload.pptx')
            if not safe.lower().endswith('.pptx'):
                self._json({'error': '仅支持 .pptx', 'type': 'unsupported_file'}, 400); return
            out_path = os.path.join(UPLOAD_DIR, f"upload_{int(time.time())}_{safe}")
            with open(out_path, 'wb') as f:
                f.write(content)
            parsed = _parse_pptx(out_path, safe)
            # 覆盖 slides.json 前先备份, 避免轻量解析覆盖原有精细 OCR 数据后不可恢复
            old_path = _slides_path()
            if os.path.exists(old_path):
                bak_path = os.path.join(WORKSPACE, f"slides_backup_{int(time.time())}.json")
                try:
                    with open(old_path, 'rb') as rf, open(bak_path, 'wb') as wf:
                        wf.write(rf.read())
                except Exception:
                    pass
            _save_json(_slides_path(), parsed)
            _SLIDES_CACHE = parsed
            _CURRENT_SOURCE = safe
            slides = parsed.get('slides', [])
            self._json({
                'ok': True,
                'source': safe,
                'page_count': len(slides),
                'n_formula_objects': _formula_count(slides),
                'n_dual_channel': _dual_count(slides),
                'cequ': '轻量解析',
                'note': '已完成轻量解析：可查看页文本、生成课件驱动题；公式 OCR/MathType 精细还原未启用',
            })
        except Exception as e:
            self._json({'error': str(e), 'type': type(e).__name__}, 500)

    def _json(self, data, status=200):
        body = json.dumps(data, ensure_ascii=False).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Cache-Control', 'no-store')
        self.send_header('Access-Control-Allow-Origin', '*')
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, fmt, *args):
        pass


def main():
    port = int(sys.argv[1]) if len(sys.argv) > 1 else 8000
    os.chdir(WORKSPACE)
    httpd = ThreadingHTTPServer(('0.0.0.0', port), Handler)
    print(f'server.py 启动: http://127.0.0.1:{port}/')
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == '__main__':
    main()
